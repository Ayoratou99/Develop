"""Chargement de documents en texte brut (Markdown, texte, PDF, CSV, JSON, YAML, HTML, DOCX, XML, SQL).

Principes :

* :func:`load_document` ne lève **jamais** pour un fichier existant : un fichier illisible,
  binaire, trop volumineux ou mal encodé renvoie un :class:`LoadedDocument` avec un texte
  vide (ou partiel) et des avertissements en français destinés à l'interface.
* Les dépôts et documents de l'utilisateur sont en lecture seule : ce module ne fait que lire.
* Aucune dépendance obligatoire au-delà de la bibliothèque standard et de ``pypdf``
  (``python-docx`` et ``pyyaml`` sont utilisés s'ils sont présents, sinon un repli stdlib).
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# Extension (sans point, minuscule) -> format logique. ``log`` est volontairement exclu.
SUPPORTED_FORMATS: Dict[str, str] = {
    "md": "markdown",
    "markdown": "markdown",
    "txt": "text",
    "text": "text",
    "rst": "text",
    "pdf": "pdf",
    "sql": "sql",
    "csv": "csv",
    "tsv": "csv",
    "json": "json",
    "yaml": "yaml",
    "yml": "yaml",
    "html": "html",
    "htm": "html",
    "docx": "docx",
    "xml": "xml",
}

CSV_MAX_ROWS = 5000  # lignes de données rendues au maximum (au-delà : avertissement)
PAGE_SEPARATOR = "\f"  # séparateur entre pages d'un PDF dans ``LoadedDocument.text``
_TITLE_MAX = 200


@dataclass
class LoadedDocument:
    """Résultat du chargement d'un document."""

    text: str
    format: str
    title: Optional[str] = None
    pages: Optional[int] = None
    warnings: List[str] = field(default_factory=list)
    meta: Dict[str, Any] = field(default_factory=dict)


@dataclass
class Table:
    """Tableau Markdown extrait d'un document (lignes 1-based inclusives, 0 si inconnues)."""

    headers: List[str]
    rows: List[List[str]]
    heading_path: List[str] = field(default_factory=list)
    line_start: int = 0
    line_end: int = 0


# --------------------------------------------------------------------------- utilitaires

def format_for_path(path: str) -> Optional[str]:
    """Format logique d'après l'extension du fichier, ou None si non pris en charge."""
    ext = Path(path).suffix.lower().lstrip(".")
    return SUPPORTED_FORMATS.get(ext)


def decode_bytes(data: bytes) -> Tuple[str, str]:
    """Décode des octets en texte avec repli d'encodage.

    Ordre : BOM UTF-8, BOM UTF-16, UTF-8 strict, puis Latin-1 (qui ne peut pas échouer).
    Renvoie ``(texte, nom_encodage)``.
    """
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig", "replace"), "utf-8-sig"
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return data.decode("latin-1"), "latin-1"


def _looks_binary(data: bytes) -> bool:
    sample = data[:8192]
    if not sample:
        return False
    if sample.startswith(b"\xff\xfe") or sample.startswith(b"\xfe\xff"):
        return False  # UTF-16 avec BOM : les NUL sont normaux
    return b"\x00" in sample


def _first_nonblank_line(text: str) -> Optional[str]:
    for line in text.splitlines():
        s = line.strip()
        if s:
            return s[:_TITLE_MAX]
    return None


def _strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


# --------------------------------------------------------------------------- Markdown

_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_FENCE_RE = re.compile(r"^\s*(```|~~~)")
_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-+:?\s*(?:\|\s*:?-+:?\s*)*\|?\s*$")


def _is_table_row(line: str) -> bool:
    s = line.strip()
    return bool(s) and "|" in s


def _is_table_separator(line: str) -> bool:
    return "|" in line and bool(_TABLE_SEP_RE.match(line))


def _split_table_cells(line: str) -> List[str]:
    s = line.strip().replace("\\|", "\x00")
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.replace("\x00", "|").strip() for c in s.split("|")]


def _iter_markdown_lines(lines: List[str]):
    """Itère ``(index, ligne, dans_une_clôture_de_code)`` en suivant les blocs ``` / ~~~."""
    fence: Optional[str] = None
    for i, line in enumerate(lines):
        m = _FENCE_RE.match(line)
        if m:
            marker = m.group(1)
            if fence is None:
                fence = marker
                yield i, line, True
                continue
            if marker == fence:
                fence = None
                yield i, line, True
                continue
        yield i, line, fence is not None


def markdown_headings(text: str) -> List[Tuple[int, str, int]]:
    """Titres ATX ``(niveau, titre, numéro de ligne 1-based)`` hors blocs de code."""
    out: List[Tuple[int, str, int]] = []
    for i, line, in_code in _iter_markdown_lines(text.splitlines()):
        if in_code:
            continue
        m = _HEADING_RE.match(line)
        if m:
            out.append((len(m.group(1)), m.group(2).strip(), i + 1))
    return out


def extract_markdown_tables(text: str) -> List[Table]:
    """Extrait les tableaux Markdown (GFM) avec le chemin des titres parents."""
    lines = text.splitlines()
    tables: List[Table] = []
    stack: List[Tuple[int, str]] = []
    rows_iter = list(_iter_markdown_lines(lines))
    i = 0
    n = len(rows_iter)
    while i < n:
        idx, line, in_code = rows_iter[i]
        if in_code:
            i += 1
            continue
        m = _HEADING_RE.match(line)
        if m:
            level = len(m.group(1))
            while stack and stack[-1][0] >= level:
                stack.pop()
            stack.append((level, m.group(2).strip()))
            i += 1
            continue
        if _is_table_row(line) and i + 1 < n and not rows_iter[i + 1][2] and _is_table_separator(rows_iter[i + 1][1]):
            headers = _split_table_cells(line)
            rows: List[List[str]] = []
            j = i + 2
            while j < n and not rows_iter[j][2] and _is_table_row(rows_iter[j][1]):
                cells = _split_table_cells(rows_iter[j][1])
                if len(cells) < len(headers):
                    cells = cells + [""] * (len(headers) - len(cells))
                rows.append(cells)
                j += 1
            tables.append(Table(headers=headers, rows=rows, heading_path=[h for _, h in stack],
                                line_start=idx + 1, line_end=rows_iter[j - 1][0] + 1))
            i = j
            continue
        i += 1
    return tables


_DD_COLUMN_WORDS = ("colonne", "column", "champ", "field", "attribut", "attribute", "propriete", "property", "nom du champ", "variable")
_DD_TYPE_WORDS = ("type", "format", "datatype", "data type")
_DD_DESC_WORDS = ("description", "libelle", "signification", "commentaire", "comment", "definition", "detail", "meaning", "remarque")


def _normalize_header(h: str) -> str:
    return _strip_accents(h).lower().strip().strip("*_`")


def looks_like_data_dictionary(text: str) -> bool:
    """Vrai si le document contient un tableau Markdown de type dictionnaire de données.

    Critère : des en-têtes du genre *colonne / champ* **et** *type* ou *description*
    (insensible à la casse et aux accents, français ou anglais), avec au moins une ligne.
    """
    for table in extract_markdown_tables(text):
        if not table.rows:
            continue
        headers = [_normalize_header(h) for h in table.headers]
        has_column = any(any(w in h for w in _DD_COLUMN_WORDS) for h in headers)
        has_type = any(any(w in h for w in _DD_TYPE_WORDS) for h in headers)
        has_desc = any(any(w in h for w in _DD_DESC_WORDS) for h in headers)
        if has_column and (has_type or has_desc):
            return True
    return False


def _markdown_title(text: str) -> Optional[str]:
    for level, title, _ in markdown_headings(text):
        if level == 1 and title:
            return title[:_TITLE_MAX]
    for level, title, _ in markdown_headings(text):
        if title:
            return title[:_TITLE_MAX]
    return _first_nonblank_line(text)


# --------------------------------------------------------------------------- PDF

def page_for_offset(page_offsets: List[int], offset: int) -> int:
    """Numéro de page (1-based) contenant la position ``offset`` dans le texte d'un PDF."""
    if not page_offsets:
        return 1
    page = 1
    for i, start in enumerate(page_offsets):
        if offset >= start:
            page = i + 1
        else:
            break
    return page


def _load_pdf(data: bytes, doc: LoadedDocument) -> None:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - pypdf est une dépendance déclarée
        doc.warnings.append("pypdf n'est pas installé : impossible de lire le PDF.")
        return
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                ok = reader.decrypt("")
            except Exception:  # noqa: BLE001 - algorithme non supporté, etc.
                ok = 0
            if not ok:
                doc.warnings.append("PDF chiffré : impossible d'extraire le texte sans mot de passe.")
                return
        n_pages = len(reader.pages)
    except Exception as exc:  # noqa: BLE001 - pypdf lève des erreurs variées sur PDF corrompu
        doc.warnings.append("PDF illisible ou corrompu : %s" % exc)
        return

    # Chaque page se termine par "\n" avant le séparateur "\f" : ``splitlines()`` produit ainsi une
    # ligne vide entre deux pages, ce qui en fait une frontière de paragraphe pour le découpage.
    parts: List[str] = []
    offsets: List[int] = []
    empty_pages = 0
    pos = 0
    for i in range(n_pages):
        try:
            page_text = reader.pages[i].extract_text() or ""
        except Exception:  # noqa: BLE001
            page_text = ""
        page_text = page_text.strip("\n") + "\n"
        if not page_text.strip():
            empty_pages += 1
        offsets.append(pos)
        parts.append(page_text)
        pos += len(page_text) + len(PAGE_SEPARATOR)
    doc.text = PAGE_SEPARATOR.join(parts)
    doc.pages = n_pages
    doc.meta["page_offsets"] = offsets
    doc.meta["empty_pages"] = empty_pages
    if n_pages and empty_pages == n_pages:
        doc.warnings.append("Aucun texte extractible : PDF scanné sans OCR ou pages vides.")
    elif empty_pages:
        doc.warnings.append("%d page(s) sans texte extractible (scan sans OCR ?)." % empty_pages)
    title = None
    try:
        md = reader.metadata
        if md is not None and md.title:
            title = str(md.title).strip() or None
    except Exception:  # noqa: BLE001
        title = None
    doc.title = title or _first_nonblank_line(doc.text)


# --------------------------------------------------------------------------- CSV

def _load_csv(text: str, doc: LoadedDocument, ext: str) -> None:
    sample = text[:20000]
    delimiter = "\t" if ext == "tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    doc.meta["delimiter"] = delimiter
    try:
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        header = next(reader, None)
        if header is None:
            doc.text = ""
            doc.meta["columns"] = []
            doc.meta["rows"] = 0
            return
        header = [h.strip() for h in header]
        out_lines = [" | ".join(header)]
        count = 0
        truncated = False
        for row in reader:
            if not any(c.strip() for c in row):
                continue
            if count >= CSV_MAX_ROWS:
                truncated = True
                break
            cells = []
            for i, val in enumerate(row):
                col = header[i] if i < len(header) and header[i] else "col%d" % (i + 1)
                cells.append("%s: %s" % (col, val.strip()))
            out_lines.append(" | ".join(cells))
            count += 1
    except csv.Error as exc:
        doc.warnings.append("CSV mal formé (%s) : contenu conservé en texte brut." % exc)
        doc.text = text
        return
    doc.text = "\n".join(out_lines)
    doc.meta["columns"] = header
    doc.meta["rows"] = count
    if truncated:
        doc.meta["truncated"] = True
        doc.warnings.append("CSV tronqué : seules les %d premières lignes ont été indexées." % CSV_MAX_ROWS)


# --------------------------------------------------------------------------- JSON / YAML / XML

def _load_json(text: str, doc: LoadedDocument) -> None:
    try:
        obj = json.loads(text)
    except ValueError as exc:
        doc.warnings.append("JSON invalide (%s) : contenu conservé en texte brut." % exc)
        doc.text = text
        doc.meta["json_valid"] = False
        return
    doc.text = json.dumps(obj, indent=2, ensure_ascii=False)
    doc.meta["json_valid"] = True
    if isinstance(obj, dict):
        doc.meta["keys"] = [str(k) for k in list(obj.keys())[:200]]


def _load_yaml(text: str, doc: LoadedDocument) -> None:
    doc.text = text
    try:
        import yaml  # type: ignore
    except ImportError:
        return
    try:
        obj = yaml.safe_load(text)
    except Exception as exc:  # noqa: BLE001 - yaml.YAMLError et dérivés
        doc.warnings.append("YAML invalide : %s" % str(exc).splitlines()[0])
        doc.meta["yaml_valid"] = False
        return
    doc.meta["yaml_valid"] = True
    if isinstance(obj, dict):
        doc.meta["keys"] = [str(k) for k in list(obj.keys())[:200]]


def _load_xml(text: str, doc: LoadedDocument) -> None:
    doc.text = text
    try:
        root = ET.fromstring(text.encode("utf-8"))
        tag = root.tag.split("}", 1)[-1]
        doc.meta["root_tag"] = tag
    except ET.ParseError as exc:
        doc.warnings.append("XML mal formé : %s" % exc)


# --------------------------------------------------------------------------- HTML

class _HTMLTextExtractor(HTMLParser):
    _SKIP = {"script", "style", "noscript", "template"}
    _BLOCK = {"p", "div", "br", "li", "ul", "ol", "h1", "h2", "h3", "h4", "h5", "h6", "tr", "table", "thead",
              "tbody", "section", "article", "header", "footer", "nav", "blockquote", "pre", "hr", "dt", "dd",
              "main", "aside", "form", "fieldset", "legend", "figure", "figcaption", "address", "option"}
    _CELL = {"td", "th"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: List[str] = []
        self.title_parts: List[str] = []
        self._skip_depth = 0
        self._in_title = False

    def handle_starttag(self, tag: str, attrs) -> None:  # type: ignore[override]
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "title":
            self._in_title = True
        elif tag in self._BLOCK:
            self.parts.append("\n")
        elif tag in self._CELL:
            self.parts.append(" | ")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP:
            self._skip_depth = max(0, self._skip_depth - 1)
        elif tag == "title":
            self._in_title = False
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        if self._in_title:
            self.title_parts.append(data)
        else:
            self.parts.append(data)


def html_to_text(html: str) -> Tuple[str, Optional[str]]:
    """Texte visible d'un document HTML (scripts/styles ignorés) et titre ``<title>``."""
    parser = _HTMLTextExtractor()
    try:
        parser.feed(html)
        parser.close()
        raw = "".join(parser.parts)
        title = " ".join("".join(parser.title_parts).split()) or None
    except Exception:  # noqa: BLE001 - HTML pathologique : repli par expression régulière
        raw = re.sub(r"<(script|style)\b.*?</\1\s*>", " ", html, flags=re.S | re.I)
        raw = re.sub(r"<[^>]+>", "\n", raw)
        title = None
    lines = [" ".join(l.split()) for l in raw.split("\n")]
    out: List[str] = []
    blank = False
    for line in lines:
        if line:
            out.append(line)
            blank = False
        elif not blank and out:
            out.append("")
            blank = True
    while out and not out[-1]:
        out.pop()
    return "\n".join(out), title


# --------------------------------------------------------------------------- DOCX

_W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_DC_NS = "{http://purl.org/dc/elements/1.1/}"


def _docx_paragraph_text(p: ET.Element) -> str:
    parts: List[str] = []
    for node in p.iter():
        if node.tag == _W_NS + "t":
            parts.append(node.text or "")
        elif node.tag == _W_NS + "tab":
            parts.append("\t")
        elif node.tag in (_W_NS + "br", _W_NS + "cr"):
            parts.append("\n")
    return "".join(parts)


def _docx_walk(elem: ET.Element, out: List[str]) -> None:
    for child in list(elem):
        if child.tag == _W_NS + "p":
            out.append(_docx_paragraph_text(child))
        elif child.tag == _W_NS + "tbl":
            for tr in child.iter(_W_NS + "tr"):
                cells = []
                for tc in tr.findall(_W_NS + "tc"):
                    cells.append(" ".join(_docx_paragraph_text(p).strip() for p in tc.iter(_W_NS + "p")).strip())
                out.append(" | ".join(cells))
            out.append("")
        else:
            _docx_walk(child, out)


def _load_docx_stdlib(data: bytes, doc: LoadedDocument) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml")
            core = z.read("docProps/core.xml") if "docProps/core.xml" in z.namelist() else b""
    except (zipfile.BadZipFile, KeyError, OSError, RuntimeError) as exc:
        doc.warnings.append("DOCX illisible (archive invalide ou sans word/document.xml) : %s" % exc)
        return
    try:
        root = ET.fromstring(xml)
    except ET.ParseError as exc:
        doc.warnings.append("DOCX illisible (XML mal formé) : %s" % exc)
        return
    lines: List[str] = []
    _docx_walk(root, lines)
    doc.text = "\n".join(lines).strip("\n")
    if core:
        try:
            title_el = ET.fromstring(core).find(".//" + _DC_NS + "title")
            if title_el is not None and (title_el.text or "").strip():
                doc.title = title_el.text.strip()[:_TITLE_MAX]
        except ET.ParseError:
            pass


def _load_docx(data: bytes, doc: LoadedDocument) -> None:
    try:
        import docx  # type: ignore
    except ImportError:
        docx = None
    if docx is not None:
        try:
            d = docx.Document(io.BytesIO(data))
            lines = [p.text for p in d.paragraphs]
            for table in d.tables:
                lines.append("")
                for row in table.rows:
                    lines.append(" | ".join(c.text.strip() for c in row.cells))
            doc.text = "\n".join(lines).strip("\n")
            title = (d.core_properties.title or "").strip()
            doc.title = title[:_TITLE_MAX] if title else None
            doc.meta["docx_backend"] = "python-docx"
            if doc.title is None:
                doc.title = _first_nonblank_line(doc.text)
            return
        except Exception:  # noqa: BLE001 - on retombe sur la lecture stdlib
            doc.text = ""
            doc.title = None
    doc.meta["docx_backend"] = "stdlib"
    _load_docx_stdlib(data, doc)
    if doc.title is None:
        doc.title = _first_nonblank_line(doc.text)


# --------------------------------------------------------------------------- SQL

_SQL_TABLE_RE = re.compile(r"create\s+(?:or\s+replace\s+)?(?:temp(?:orary)?\s+)?table\s+(?:if\s+not\s+exists\s+)?[`\"\[]?([\w.]+)",
                           re.I)


def _load_sql(text: str, doc: LoadedDocument, name: str) -> None:
    doc.text = text
    doc.title = name
    tables = []
    for m in _SQL_TABLE_RE.finditer(text):
        t = m.group(1).split(".")[-1]
        if t not in tables:
            tables.append(t)
    doc.meta["tables"] = tables


# --------------------------------------------------------------------------- point d'entrée

def load_document(path: str, max_bytes: int = 20_000_000, *, format_hint: Optional[str] = None) -> LoadedDocument:
    """Charge un fichier en texte selon son extension (voir :data:`SUPPORTED_FORMATS`).

    Ne lève jamais pour un fichier existant : en cas de problème, ``text`` est vide (ou partiel)
    et ``warnings`` explique pourquoi. ``format_hint`` force le format (ex: ``"markdown"``).
    """
    p = Path(path)
    ext = p.suffix.lower().lstrip(".")
    fmt = format_hint or SUPPORTED_FORMATS.get(ext)
    doc = LoadedDocument(text="", format=fmt or "unknown", meta={"path": str(p), "extension": ext})

    if fmt is None:
        doc.warnings.append("Format non pris en charge : .%s (%s)." % (ext or "?", p.name))
        return doc
    if fmt not in set(SUPPORTED_FORMATS.values()):
        doc.warnings.append("Format inconnu : %s." % fmt)
        doc.format = "unknown"
        return doc
    try:
        if not p.is_file():
            doc.warnings.append("Fichier introuvable ou non régulier : %s" % p)
            return doc
        size = p.stat().st_size
    except OSError as exc:
        doc.warnings.append("Fichier inaccessible : %s" % exc)
        return doc
    doc.meta["size"] = size
    if size > max_bytes:
        doc.warnings.append("Fichier trop volumineux (%d octets > %d) : non indexé." % (size, max_bytes))
        return doc
    try:
        data = p.read_bytes()
    except OSError as exc:
        doc.warnings.append("Lecture impossible : %s" % exc)
        return doc

    if fmt == "pdf":
        _load_pdf(data, doc)
        return doc
    if fmt == "docx":
        _load_docx(data, doc)
        return doc

    if _looks_binary(data):
        doc.warnings.append("Fichier binaire ou encodage inconnu : contenu ignoré (%s)." % p.name)
        return doc
    text, encoding = decode_bytes(data)
    doc.meta["encoding"] = encoding
    if encoding == "latin-1":
        doc.warnings.append("Encodage non UTF-8 : lecture en Latin-1 (certains caractères peuvent être altérés).")
    text = text.replace("\r\n", "\n").replace("\r", "\n")

    if fmt == "markdown":
        doc.text = text
        doc.title = _markdown_title(text)
        doc.meta["headings"] = [t for _, t, _ in markdown_headings(text)]
        doc.meta["looks_like_data_dictionary"] = looks_like_data_dictionary(text)
    elif fmt == "text":
        doc.text = text
        doc.title = _first_nonblank_line(text)
    elif fmt == "csv":
        _load_csv(text, doc, ext)
        doc.title = p.name
    elif fmt == "json":
        _load_json(text, doc)
        doc.title = p.name
    elif fmt == "yaml":
        _load_yaml(text, doc)
        doc.title = p.name
    elif fmt == "html":
        doc.text, title = html_to_text(text)
        doc.title = title or _first_nonblank_line(doc.text)
    elif fmt == "xml":
        _load_xml(text, doc)
        doc.title = p.name
    elif fmt == "sql":
        _load_sql(text, doc, p.name)

    if doc.text and fmt in ("text", "html"):
        doc.meta["looks_like_data_dictionary"] = looks_like_data_dictionary(doc.text)
    if not doc.text.strip() and not doc.warnings:
        doc.warnings.append("Document vide : aucun texte à indexer (%s)." % p.name)
    return doc
